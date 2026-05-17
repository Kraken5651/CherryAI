"""
Study Mode Router - Document processing, Q&A, flashcards, quizzes, summaries.
"""
import os
import json
import uuid
import shutil
from pathlib import Path
from fastapi import APIRouter, HTTPException, UploadFile, File, Form
from pydantic import BaseModel
from google.genai import types
from gemini_client import generate_with_fallback
from pypdf import PdfReader
from pptx import Presentation

router = APIRouter()

# Storage directory for uploaded documents
UPLOAD_DIR = Path("uploads/study")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

# In-memory document store (text content indexed by doc_id)
documents: dict[str, dict] = {}


def extract_text_from_pdf(file_path: Path) -> str:
    reader = PdfReader(str(file_path))
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n\n"
    return text.strip()


def extract_text_from_pptx(file_path: Path) -> str:
    prs = Presentation(str(file_path))
    text = ""
    for slide_num, slide in enumerate(prs.slides, 1):
        text += f"--- Slide {slide_num} ---\n"
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text += shape.text + "\n"
        text += "\n"
    return text.strip()


def extract_text_from_txt(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8", errors="ignore")


EXTRACTORS = {
    ".pdf": extract_text_from_pdf,
    ".pptx": extract_text_from_pptx,
    ".ppt": extract_text_from_pptx,
    ".txt": extract_text_from_txt,
    ".md": extract_text_from_txt,
}


# ── Upload ────────────────────────────────────────────────────

@router.post("/upload")
async def upload_document(file: UploadFile = File(...)):
    ext = Path(file.filename or "").suffix.lower()
    if ext not in EXTRACTORS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type: {ext}. Supported: {', '.join(EXTRACTORS.keys())}"
        )

    doc_id = uuid.uuid4().hex[:12]
    save_path = UPLOAD_DIR / f"{doc_id}{ext}"

    with open(save_path, "wb") as f:
        shutil.copyfileobj(file.file, f)

    try:
        text = EXTRACTORS[ext](save_path)
    except Exception as e:
        save_path.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail=f"Failed to extract text: {e}")

    if not text.strip():
        save_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="No text could be extracted from this file.")

    documents[doc_id] = {
        "id": doc_id,
        "filename": file.filename,
        "ext": ext,
        "text": text,
        "path": str(save_path),
        "char_count": len(text),
        "page_count": len(PdfReader(str(save_path)).pages) if ext == ".pdf" else None,
    }

    return {
        "id": doc_id,
        "filename": file.filename,
        "char_count": len(text),
        "page_count": documents[doc_id]["page_count"],
        "message": "Document uploaded and processed successfully.",
    }


# ── List Documents ────────────────────────────────────────────

@router.get("/documents")
async def list_documents():
    return [
        {
            "id": doc["id"],
            "filename": doc["filename"],
            "char_count": doc["char_count"],
            "page_count": doc["page_count"],
        }
        for doc in documents.values()
    ]


# ── Delete Document ───────────────────────────────────────────

@router.delete("/documents/{doc_id}")
async def delete_document(doc_id: str):
    if doc_id not in documents:
        raise HTTPException(status_code=404, detail="Document not found.")
    
    Path(documents[doc_id]["path"]).unlink(missing_ok=True)
    del documents[doc_id]
    return {"message": "Document deleted."}


# ── Ask Questions ─────────────────────────────────────────────

class AskRequest(BaseModel):
    doc_id: str
    question: str

@router.post("/ask")
async def ask_question(request: AskRequest):
    if request.doc_id not in documents:
        raise HTTPException(status_code=404, detail="Document not found.")

    doc = documents[request.doc_id]
    # Truncate text to ~60k chars to stay within token limits
    context = doc["text"][:60000]

    response = generate_with_fallback(
        contents=[
            types.Content(role="user", parts=[
                types.Part.from_text(text=f"""You are Cherry AI, a study assistant. Answer the following question based ONLY on the provided document content. 
If the answer is not in the document, say so clearly.

DOCUMENT ({doc['filename']}):
\"\"\"
{context}
\"\"\"

QUESTION: {request.question}

Answer clearly and concisely:""")
            ])
        ]
    )

    return {"answer": response.text}


# ── Summarize ─────────────────────────────────────────────────

class SummarizeRequest(BaseModel):
    doc_id: str

@router.post("/summarize")
async def summarize_document(request: SummarizeRequest):
    if request.doc_id not in documents:
        raise HTTPException(status_code=404, detail="Document not found.")

    doc = documents[request.doc_id]
    context = doc["text"][:60000]

    response = generate_with_fallback(
        contents=[
            types.Content(role="user", parts=[
                types.Part.from_text(text=f"""Summarize the following document clearly and concisely. 
Organize with headings and bullet points. Highlight key concepts.

DOCUMENT ({doc['filename']}):
\"\"\"
{context}
\"\"\"

Provide a comprehensive summary:""")
            ])
        ]
    )

    return {"summary": response.text}


# ── Generate Flashcards ──────────────────────────────────────

class FlashcardRequest(BaseModel):
    doc_id: str
    count: int = 10

@router.post("/flashcards")
async def generate_flashcards(request: FlashcardRequest):
    if request.doc_id not in documents:
        raise HTTPException(status_code=404, detail="Document not found.")

    doc = documents[request.doc_id]
    context = doc["text"][:60000]

    response = generate_with_fallback(
        contents=[
            types.Content(role="user", parts=[
                types.Part.from_text(text=f"""Generate exactly {request.count} flashcards from this document for study purposes.
Return ONLY a valid JSON array. Each item must have "front" (question) and "back" (answer) fields.
Keep answers concise but complete.

DOCUMENT ({doc['filename']}):
\"\"\"
{context}
\"\"\"

Return JSON array:""")
            ])
        ]
    )

    # Try to parse JSON from response
    text = response.text.strip()
    # Remove markdown code fences if present
    if text.startswith("```"):
        text = text.split("\n", 1)[-1]
        text = text.rsplit("```", 1)[0]

    try:
        flashcards = json.loads(text)
    except json.JSONDecodeError:
        flashcards = [{"front": "Error parsing flashcards", "back": text}]

    return {"flashcards": flashcards}


# ── Generate Quiz ─────────────────────────────────────────────

class QuizRequest(BaseModel):
    doc_id: str
    count: int = 5

@router.post("/quiz")
async def generate_quiz(request: QuizRequest):
    if request.doc_id not in documents:
        raise HTTPException(status_code=404, detail="Document not found.")

    doc = documents[request.doc_id]
    context = doc["text"][:60000]

    response = generate_with_fallback(
        contents=[
            types.Content(role="user", parts=[
                types.Part.from_text(text=f"""Generate exactly {request.count} multiple-choice questions from this document.
Return ONLY a valid JSON array. Each item must have:
- "question": the question text
- "options": array of 4 options (strings)
- "correct": index of the correct option (0-3)
- "explanation": brief explanation of the answer

DOCUMENT ({doc['filename']}):
\"\"\"
{context}
\"\"\"

Return JSON array:""")
            ])
        ]
    )

    text = response.text.strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[-1]
        text = text.rsplit("```", 1)[0]

    try:
        quiz = json.loads(text)
    except json.JSONDecodeError:
        quiz = [{"question": "Error parsing quiz", "options": ["N/A"], "correct": 0, "explanation": text}]

    return {"quiz": quiz}
